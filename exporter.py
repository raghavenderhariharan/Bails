"""Build an .xlsx export that mirrors the layout of BailsLedgerBook.xlsx.

Two sheets reproduce the original workbook (``Dashboard`` and
``Transaction Ledger``); a third, ``Partner Profit``, adds the equity split the
app computes. Regenerating always produces the same shape, so the file can be
overwritten in place as "the latest".
"""
from datetime import date, timedelta, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from finance import (
    apply_period, balance_as_of, opening_balance, q2, totals,
)
from models import EXPENSE, INCOME, Partner, Transaction

# Brand palette, taken from the ground's badge.
GREEN_DARK = "17593C"
GREEN = "2CAA70"
CREAM = "FDF4A3"
RED = "D72228"
NAVY = "273076"
GREY_BG = "F1F4EF"

MONEY = '#,##0.00'
MONEY_INT = '#,##0'
DATE_FMT = 'dd mmm yyyy'
MONTH_FMT = 'mmm yyyy'

# Mirrors the "Recurring Commitment" block in the original workbook. Edit here
# if the standing commitments change.
RECURRING_COMMITMENTS = [
    ("Market allowance — Groundsman 1", "Weekly", 1000),
    ("Market allowance — Groundsman 2", "Weekly", 1000),
    ("Salary / chit payment — Groundsman 1", "Every 15 days", 5000),
]

THIN = Side(style="thin", color="D3DAD0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _title(ws, cell_range, text):
    first = cell_range.split(":")[0]
    ws.merge_cells(cell_range)
    cell = ws[first]
    cell.value = text
    cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _tile(ws, label_range, value_range, label, value, colour):
    first_label = label_range.split(":")[0]
    ws.merge_cells(label_range)
    cell = ws[first_label]
    cell.value = label
    cell.font = Font(bold=True, size=9, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=colour)
    cell.alignment = Alignment(horizontal="center", vertical="center")

    first_value = value_range.split(":")[0]
    ws.merge_cells(value_range)
    cell = ws[first_value]
    cell.value = float(value)
    cell.number_format = MONEY_INT
    cell.font = Font(bold=True, size=18, color=colour)
    cell.fill = PatternFill("solid", fgColor=GREY_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _header_row(ws, row, headers, start=1, colour=GREEN_DARK):
    for offset, text in enumerate(headers):
        cell = ws.cell(row=row, column=start + offset, value=text)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=colour)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def _widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _build_dashboard(wb, period, summary, trend):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    _widths(ws, [30, 14, 4, 18, 14, 4, 16, 14])

    _title(ws, "A1:H2", "BAILS GROUND BOOKINGS — LEDGER DASHBOARD")
    ws.row_dimensions[1].height = 22

    ws["A4"] = "Selected Period"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = period.label
    ws["B4"].font = Font(italic=True)

    _tile(ws, "A6:B6", "A7:B8", "INCOME (CREDIT)", summary["income"], GREEN_DARK)
    _tile(ws, "D6:E6", "D7:E8", "EXPENSE (DEBIT)", summary["expense"], RED)
    _tile(ws, "G6:H6", "G7:H8", "NET BALANCE", summary["net"], NAVY)
    ws.row_dimensions[7].height = 26

    _header_row(ws, 11, ["Month", "Income", "Expense", "Net"])
    for offset, month in enumerate(trend):
        row = 12 + offset
        year, mon = (int(part) for part in month["key"].split("-"))
        cell = ws.cell(row=row, column=1, value=date(year, mon, 1))
        cell.number_format = MONTH_FMT
        for column, key in ((2, "income"), (3, "expense"), (4, "net")):
            value = ws.cell(row=row, column=column, value=float(month[key]))
            value.number_format = MONEY_INT
        for column in range(1, 5):
            ws.cell(row=row, column=column).border = BOX

    total_row = 12 + len(trend)
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    for column, key in ((2, "income"), (3, "expense"), (4, "net")):
        cell = ws.cell(row=total_row, column=column,
                       value=float(sum(m[key] for m in trend)))
        cell.number_format = MONEY_INT
        cell.font = Font(bold=True)
    for column in range(1, 5):
        ws.cell(row=total_row, column=column).fill = PatternFill("solid", fgColor=CREAM)
        ws.cell(row=total_row, column=column).border = BOX

    start = total_row + 3
    ws.cell(row=start - 1, column=1, value="Standing commitments").font = Font(bold=True, size=11)
    _header_row(ws, start, ["Recurring Commitment", "Frequency", "Amount", "Notes"])
    for offset, (name, frequency, amount) in enumerate(RECURRING_COMMITMENTS, start=1):
        row = start + offset
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=frequency)
        cell = ws.cell(row=row, column=3, value=amount)
        cell.number_format = MONEY_INT
        ws.cell(row=row, column=4, value="Recorded as an expense entry when paid")
        for column in range(1, 5):
            ws.cell(row=row, column=column).border = BOX
    return ws


def _build_ledger(wb, period):
    ws = wb.create_sheet("Transaction Ledger")
    ws.freeze_panes = "A2"
    _widths(ws, [13, 11, 10, 22, 34, 12, 20, 13, 13, 13, 16, 30])

    headers = ["Date", "Day", "Type", "Category", "Description", "Slot", "Party",
               "Income", "Expense", "Net", "Running Balance", "Notes"]
    _header_row(ws, 1, headers)

    rows = (
        apply_period(Transaction.query, period)
        .order_by(Transaction.txn_date.asc(), Transaction.id.asc())
        .all()
    )

    # Opening balance line, as in the original workbook.
    start_balance = (
        balance_as_of(period.start - timedelta(days=1)) if period.start else opening_balance()
    )
    # One date drives both the cell and its day name, so they cannot disagree.
    opening_date = period.start or (rows[0].txn_date if rows else date.today())
    ws.cell(row=2, column=1, value=opening_date).number_format = DATE_FMT
    ws.cell(row=2, column=2, value=opening_date.strftime("%A"))
    ws.cell(row=2, column=3, value="Opening Balance")
    ws.cell(row=2, column=4, value="Opening balance")
    ws.cell(row=2, column=5, value=f"Starting balance for {period.label}")
    for column, value in ((8, 0), (9, 0), (10, 0), (11, float(start_balance))):
        cell = ws.cell(row=2, column=column, value=value)
        cell.number_format = MONEY
    for column in range(1, 13):
        ws.cell(row=2, column=column).fill = PatternFill("solid", fgColor=GREY_BG)
        ws.cell(row=2, column=column).border = BOX

    running = start_balance
    for offset, txn in enumerate(rows):
        row = 3 + offset
        running = q2(running + txn.signed_amount)
        income = float(txn.amount) if txn.kind == INCOME else 0.0
        expense = float(txn.amount) if txn.kind == EXPENSE else 0.0

        ws.cell(row=row, column=1, value=txn.txn_date).number_format = DATE_FMT
        ws.cell(row=row, column=2, value=txn.day_name)
        ws.cell(row=row, column=3, value="Income" if txn.kind == INCOME else "Expense")
        ws.cell(row=row, column=4, value=txn.category)
        ws.cell(row=row, column=5, value=txn.description)
        ws.cell(row=row, column=6, value=txn.slot)
        ws.cell(row=row, column=7, value=txn.party)
        for column, value in ((8, income), (9, expense),
                              (10, float(txn.signed_amount)), (11, float(running))):
            cell = ws.cell(row=row, column=column, value=value)
            cell.number_format = MONEY
        ws.cell(row=row, column=12, value=txn.notes)

        colour = GREEN if txn.kind == INCOME else RED
        ws.cell(row=row, column=3).font = Font(bold=True, color=colour)
        ws.cell(row=row, column=10).font = Font(color=colour)
        for column in range(1, 13):
            ws.cell(row=row, column=column).border = BOX

    last = 2 + len(rows)
    total_row = last + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=f"{period.label}").font = Font(bold=True, italic=True)
    income_total = sum(float(t.amount) for t in rows if t.kind == INCOME)
    expense_total = sum(float(t.amount) for t in rows if t.kind == EXPENSE)
    for column, value in ((8, income_total), (9, expense_total),
                          (10, income_total - expense_total), (11, float(running))):
        cell = ws.cell(row=total_row, column=column, value=value)
        cell.number_format = MONEY
        cell.font = Font(bold=True)
    for column in range(1, 13):
        ws.cell(row=total_row, column=column).fill = PatternFill("solid", fgColor=CREAM)
        ws.cell(row=total_row, column=column).border = BOX

    # No auto_filter here on purpose: openpyxl expresses it as an
    # `_xlnm._FilterDatabase` defined name, which Excel frequently rejects with
    # "We found a problem with some content". The frozen header row gives the
    # same practical benefit, and the reader can switch filtering on themselves.
    return ws


def _build_partners(wb, period, summary, shares):
    ws = wb.create_sheet("Partner Profit")
    ws.sheet_view.showGridLines = False
    _widths(ws, [30, 14, 18, 18])

    _title(ws, "A1:D2", "PARTNER PROFIT SHARE")
    ws["A4"] = "Period"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = period.label

    ws["A5"] = "Distributable net profit"
    ws["A5"].font = Font(bold=True)
    cell = ws["B5"]
    cell.value = float(summary["net"])
    cell.number_format = MONEY
    cell.font = Font(bold=True, color=NAVY if summary["net"] >= 0 else RED)

    _header_row(ws, 7, ["Partner", "Equity %", "Share of net profit", "Status"])
    for offset, row in enumerate(shares["rows"], start=1):
        line = 7 + offset
        ws.cell(row=line, column=1, value=row["name"])
        equity = ws.cell(row=line, column=2, value=float(row["equity"]) / 100.0)
        equity.number_format = '0.00%'
        share = ws.cell(row=line, column=3, value=float(row["share"]))
        share.number_format = MONEY
        ws.cell(row=line, column=4,
                value="Active" if row["partner"].is_active else "Excluded")
        for column in range(1, 5):
            ws.cell(row=line, column=column).border = BOX

    total_row = 8 + len(shares["rows"])
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    equity = ws.cell(row=total_row, column=2, value=float(shares["equity_total"]) / 100.0)
    equity.number_format = '0.00%'
    equity.font = Font(bold=True)
    allocated = ws.cell(row=total_row, column=3, value=float(shares["allocated"]))
    allocated.number_format = MONEY
    allocated.font = Font(bold=True)
    for column in range(1, 5):
        ws.cell(row=total_row, column=column).fill = PatternFill("solid", fgColor=CREAM)
        ws.cell(row=total_row, column=column).border = BOX

    if not shares["balanced"]:
        note = ws.cell(row=total_row + 2, column=1,
                       value=f"Active equity totals {shares['equity_total']}%, not 100%. "
                             f"Unallocated: {shares['unallocated']}")
        note.font = Font(italic=True, color=RED)
    return ws


def build_workbook(period, generated_at=None):
    """Assemble the whole workbook for a period. Caller supplies the timestamp."""
    from finance import partner_shares, trend_window

    summary = totals(period)
    partners = Partner.query.filter_by(is_active=True).order_by(
        Partner.sort_order, Partner.id).all()
    shares = partner_shares(summary["net"], partners)
    # Same 12-month window the dashboard chart shows, so the two never disagree.
    trend = trend_window(period)["series"]

    wb = Workbook()
    wb.remove(wb.active)
    _build_dashboard(wb, period, summary, trend)
    _build_ledger(wb, period)
    _build_partners(wb, period, summary, shares)

    # An empty <workbookProtection/> element serves no purpose and is one more
    # thing for Excel to object to.
    wb.security = None

    props = wb.properties
    props.title = "Bails Cricket Ground - Ledger"
    props.creator = "Bails Ledger"
    if generated_at is not None:
        # openpyxl appends "Z" to whatever it serialises, so a tz-aware value
        # produces "...+00:00Z" -- invalid W3CDTF, and Excel then offers to
        # "recover" the file. Store a naive UTC timestamp instead.
        stamp = generated_at
        if stamp.tzinfo is not None:
            stamp = stamp.astimezone(timezone.utc).replace(tzinfo=None)
        props.created = stamp
        props.modified = stamp
    wb.active = 0
    return wb
