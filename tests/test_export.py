"""Excel export: structure mirrors BailsLedgerBook.xlsx, and the figures tie out."""
import os
import re
import sys
from datetime import date, datetime, timezone
from decimal import Decimal

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("SECRET_KEY", "test-secret")


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/x.db")
    monkeypatch.setenv("SECRET_KEY", "test-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "Admin123")
    monkeypatch.setenv("DRIVE_EXPORT_DIR", "")
    for module in ("config", "models", "finance", "exporter", "app"):
        sys.modules.pop(module, None)
    import app as app_module
    app_module._login_attempts.clear()
    app_module._global_failures.clear()
    app_module.app.config.update(TESTING=True)

    from models import EXPENSE, INCOME, Transaction, db
    rows = [
        ("2026-10-02", INCOME, "Ground booking", "7:00 AM slot", "7000"),
        ("2026-10-02", INCOME, "Ground booking", "10:45 AM slot", "6000"),
        ("2026-10-02", INCOME, "Ground booking", "2:30 PM slot", "6000"),
        ("2026-10-02", INCOME, "Other income", "Unallocated", "1000"),
        ("2026-10-03", INCOME, "Ground booking", "7:00 AM slot", "8000"),
        ("2026-10-03", INCOME, "Ground booking", "2:30 PM slot", "6000"),
        ("2026-10-04", INCOME, "ECL booking", "All three slots", "17000"),
        ("2026-10-04", EXPENSE, "Market expenses", "Groundsman 1", "1000"),
        ("2026-10-04", EXPENSE, "Market expenses", "Groundsman 2", "1000"),
        ("2026-10-04", EXPENSE, "Salary / chit payment", "Chit", "5000"),
    ]
    with app_module.app.app_context():
        for day, kind, category, description, amount in rows:
            db.session.add(Transaction(txn_date=date.fromisoformat(day), kind=kind,
                                       category=category, description=description,
                                       amount=Decimal(amount)))
        db.session.commit()
    return app_module.app


def build(app, args=None):
    from exporter import build_workbook
    from finance import parse_period
    with app.app_context():
        return build_workbook(parse_period(args or {"period": "all"}),
                              generated_at=datetime(2026, 10, 5, tzinfo=timezone.utc))


def saved(app, tmp_path, args=None):
    path = tmp_path / "out.xlsx"
    build(app, args).save(path)
    return openpyxl.load_workbook(path, data_only=True)


# ------------------------------- structure --------------------------------- #

def test_sheets_mirror_the_original_workbook(app, tmp_path):
    wb = saved(app, tmp_path)
    assert wb.sheetnames[:2] == ["Dashboard", "Transaction Ledger"]
    assert "Partner Profit" in wb.sheetnames


def test_ledger_columns_match_the_original(app, tmp_path):
    ws = saved(app, tmp_path)["Transaction Ledger"]
    header = [ws.cell(row=1, column=c).value for c in range(1, 13)]
    assert header == ["Date", "Day", "Type", "Category", "Description", "Slot",
                      "Party", "Income", "Expense", "Net", "Running Balance", "Notes"]


def test_dashboard_has_the_tiles_and_month_table(app, tmp_path):
    ws = saved(app, tmp_path)["Dashboard"]
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("LEDGER DASHBOARD" in t for t in text)
    assert "INCOME (CREDIT)" in text and "EXPENSE (DEBIT)" in text and "NET BALANCE" in text
    assert "Month" in text and "Income" in text and "Expense" in text and "Net" in text
    assert any("Recurring Commitment" in t for t in text)


def test_opening_balance_row_is_present_and_self_consistent(app, tmp_path):
    ws = saved(app, tmp_path)["Transaction Ledger"]
    assert ws.cell(row=2, column=3).value == "Opening Balance"
    day_cell = ws.cell(row=2, column=1).value
    assert day_cell.strftime("%A") == ws.cell(row=2, column=2).value


# -------------------------------- figures ---------------------------------- #

def test_totals_tie_out_to_the_ledger(app, tmp_path):
    ws = saved(app, tmp_path)["Transaction Ledger"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    total = next(r for r in rows if r[0] == "TOTAL")
    assert total[7] == 51000
    assert total[8] == 7000
    assert total[9] == 44000
    assert total[10] == 44000  # final running balance


def test_running_balance_is_cumulative(app, tmp_path):
    ws = saved(app, tmp_path)["Transaction Ledger"]
    balances = [ws.cell(row=r, column=11).value for r in range(3, 13)]
    assert balances == [7000, 13000, 19000, 20000, 28000, 34000, 51000, 50000, 49000, 44000]


def test_dashboard_tiles_match_the_data(app, tmp_path):
    ws = saved(app, tmp_path)["Dashboard"]
    assert ws["A7"].value == 51000
    assert ws["D7"].value == 7000
    assert ws["G7"].value == 44000


def test_partner_sheet_splits_the_net_exactly(app, tmp_path):
    ws = saved(app, tmp_path)["Partner Profit"]
    rows = [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value,
             ws.cell(row=r, column=3).value) for r in range(8, 14)]
    assert [r[2] for r in rows] == [8800, 8800, 8800, 8800, 4400, 4400]
    assert sum(r[2] for r in rows) == 44000
    # Equity is written as a real percentage, so Excel formats it as 20.00%.
    assert [r[1] for r in rows] == [0.2, 0.2, 0.2, 0.2, 0.1, 0.1]
    total = ws.cell(row=14, column=3).value
    assert total == 44000


def test_period_filter_narrows_the_export(app, tmp_path):
    ws = saved(app, tmp_path, {"period": "custom", "from": "2026-10-01", "to": "2026-10-03"})["Transaction Ledger"]
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] != "TOTAL"]
    assert len(rows) == 6
    total = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == "TOTAL")
    assert total[7] == 34000 and total[8] == 0


def test_export_survives_an_empty_period(app, tmp_path):
    wb = saved(app, tmp_path, {"period": "month", "month": "2027-05"})
    ws = wb["Transaction Ledger"]
    assert ws.cell(row=2, column=3).value == "Opening Balance"
    assert wb["Dashboard"]["A7"].value == 0


# --------------------------------- routes ---------------------------------- #

def _csrf(html):
    return re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)


def admin(app):
    client = app.test_client()
    page = client.get("/login?as=admin").get_data(as_text=True)
    client.post("/login", data={"csrf_token": _csrf(page), "password": "Admin123"})
    return client


def partner(app):
    client = app.test_client()
    page = client.get("/login").get_data(as_text=True)
    client.post("/login/partner", data={"csrf_token": _csrf(page)})
    return client


def test_xlsx_download_works_for_both_roles(app):
    for client in (admin(app), partner(app)):
        response = client.get("/export.xlsx?period=all")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["Content-Type"]
        assert ".xlsx" in response.headers["Content-Disposition"]
        body = response.get_data()
        assert body[:2] == b"PK", "not a real xlsx (zip) payload"
        assert len(body) > 5000


def test_xlsx_download_requires_a_login(app):
    response = app.test_client().get("/export.xlsx")
    assert response.status_code in (301, 302, 308)


def test_drive_button_is_hidden_when_no_folder_is_configured(app):
    body = admin(app).get("/").get_data(as_text=True)
    assert "Save to Drive" not in body
    assert "Excel" in body  # the download is always offered


def test_drive_sync_writes_the_file(tmp_path, monkeypatch):
    target = tmp_path / "drive"
    target.mkdir()
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/d.db")
    monkeypatch.setenv("SECRET_KEY", "test-secret")
    monkeypatch.setenv("ADMIN_PASSWORD", "Admin123")
    monkeypatch.setenv("DRIVE_EXPORT_DIR", str(target))
    monkeypatch.setenv("DRIVE_EXPORT_FILENAME", "BailsLedgerBook.xlsx")
    for module in ("config", "models", "finance", "exporter", "app"):
        sys.modules.pop(module, None)
    import app as app_module
    app_module._login_attempts.clear()
    app_module._global_failures.clear()
    app_module.app.config.update(TESTING=True)

    client = admin(app_module.app)
    body = client.get("/").get_data(as_text=True)
    assert "Save to Drive" in body, "button missing although the folder exists"

    page = client.get("/").get_data(as_text=True)
    response = client.post("/export/drive?period=all",
                           data={"csrf_token": _csrf(page), "next": "/"},
                           follow_redirects=True)
    assert response.status_code == 200
    written = target / "BailsLedgerBook.xlsx"
    assert written.exists(), "nothing written to the sync folder"
    assert written.read_bytes()[:2] == b"PK"
    assert not list(target.glob("*.part")), "temp file left behind"

    # Re-running replaces the file in place rather than piling up copies.
    first = written.stat().st_size
    page = client.get("/").get_data(as_text=True)
    client.post("/export/drive?period=all", data={"csrf_token": _csrf(page), "next": "/"},
                follow_redirects=True)
    assert len(list(target.glob("*.xlsx"))) == 1
    assert written.stat().st_size == pytest.approx(first, rel=0.1)


def test_partner_cannot_push_to_drive(tmp_path, monkeypatch):
    target = tmp_path / "drive2"
    target.mkdir()
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/d2.db")
    monkeypatch.setenv("SECRET_KEY", "test-secret")
    monkeypatch.setenv("DRIVE_EXPORT_DIR", str(target))
    for module in ("config", "models", "finance", "exporter", "app"):
        sys.modules.pop(module, None)
    import app as app_module
    app_module._login_attempts.clear()
    app_module._global_failures.clear()
    app_module.app.config.update(TESTING=True)

    client = partner(app_module.app)
    page = client.get("/").get_data(as_text=True)
    assert "Save to Drive" not in page
    response = client.post("/export/drive?period=all",
                           data={"csrf_token": _csrf(page), "next": "/"})
    assert response.status_code == 403
    assert not list(target.glob("*.xlsx")), "partner wrote to the sync folder"


# ------------------- workbook validity (Excel repair prompt) ---------------- #

def test_core_properties_dates_are_valid_w3cdtf(app, tmp_path):
    """A tz-aware timestamp used to serialise as "...+00:00Z", which is invalid
    W3CDTF and made Excel offer to recover the file on open."""
    import re
    import zipfile

    path = tmp_path / "props.xlsx"
    build(app).save(path)
    core = zipfile.ZipFile(path).read("docProps/core.xml").decode()

    pattern = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")
    for tag in ("created", "modified"):
        match = re.search(rf"<dcterms:{tag}[^>]*>([^<]+)<", core)
        assert match, f"no dcterms:{tag} in core.xml"
        value = match.group(1)
        assert pattern.match(value), f"dcterms:{tag} is not valid W3CDTF: {value!r}"
        assert "+00:00" not in value, f"dcterms:{tag} carries both an offset and Z: {value!r}"


def test_every_xml_part_is_well_formed(app, tmp_path):
    import zipfile
    from xml.etree import ElementTree

    path = tmp_path / "parts.xlsx"
    build(app).save(path)
    archive = zipfile.ZipFile(path)
    checked = 0
    for name in archive.namelist():
        if name.endswith((".xml", ".rels")):
            ElementTree.fromstring(archive.read(name))  # raises on malformed XML
            checked += 1
    assert checked >= 8, "suspiciously few XML parts"


def test_no_duplicate_cell_references(app, tmp_path):
    """Writing into a merged cell's non-anchor slot produces duplicate refs,
    which Excel also reports as damage."""
    import collections
    import re
    import zipfile

    path = tmp_path / "cells.xlsx"
    build(app).save(path)
    archive = zipfile.ZipFile(path)
    for name in archive.namelist():
        if name.startswith("xl/worksheets/sheet"):
            refs = re.findall(r'<c r="([A-Z]+\d+)"', archive.read(name).decode())
            duplicates = [r for r, n in collections.Counter(refs).items() if n > 1]
            assert not duplicates, f"{name} has duplicate cell refs: {duplicates}"


def test_naive_timestamp_is_accepted_too(app, tmp_path):
    """The caller may pass a naive datetime; it must not be mangled either."""
    import re
    import zipfile

    from exporter import build_workbook
    from finance import parse_period

    path = tmp_path / "naive.xlsx"
    with app.app_context():
        build_workbook(parse_period({"period": "all"}),
                       generated_at=datetime(2026, 10, 5, 12, 30, 0)).save(path)
    core = zipfile.ZipFile(path).read("docProps/core.xml").decode()
    created = re.search(r"<dcterms:created[^>]*>([^<]+)<", core).group(1)
    assert created == "2026-10-05T12:30:00Z", created


def test_no_autofilter_defined_name(app, tmp_path):
    """openpyxl writes an autofilter as an `_xlnm._FilterDatabase` defined name,
    which Excel frequently rejects with "We found a problem with some content"."""
    import zipfile

    path = tmp_path / "nofilter.xlsx"
    build(app).save(path)
    archive = zipfile.ZipFile(path)
    workbook_xml = archive.read("xl/workbook.xml").decode()
    assert "_FilterDatabase" not in workbook_xml
    for name in archive.namelist():
        if name.startswith("xl/worksheets/sheet"):
            assert "autoFilter" not in archive.read(name).decode(), f"{name} has an autoFilter"


def test_ledger_sheet_rows_are_ascending(app, tmp_path):
    ws = saved(app, tmp_path)["Transaction Ledger"]
    dates = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row)]
    dates = [d for d in dates if hasattr(d, "year")]
    assert dates == sorted(dates), f"ledger sheet not ascending: {dates}"


def test_frozen_header_survives(app, tmp_path):
    """The frozen pane is what replaces the autofilter's usefulness."""
    import zipfile

    path = tmp_path / "freeze.xlsx"
    build(app).save(path)
    sheet = zipfile.ZipFile(path).read("xl/worksheets/sheet2.xml").decode()
    assert 'state="frozen"' in sheet
